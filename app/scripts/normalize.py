"""Versioned app transformation; never rewrites supplied sources or extractions."""
from pathlib import Path
import json, hashlib
from decimal import Decimal
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'app/data'
OUT.mkdir(exist_ok=True)
read = lambda name: json.loads((ROOT / 'processed' / f'{name}.json').read_text())
inventory = read('source_inventory')
assert len(inventory) == 22
for source in inventory:
    path = ROOT / source['path']
    assert path.is_file(), f"Missing original: {path}"
    assert hashlib.sha256(path.read_bytes()).hexdigest() == source['sha256'], f"Changed original: {path}"
manifest = json.loads((ROOT / 'PACKAGE_MANIFEST.json').read_text())
discrepancies = []
for item in manifest['files']:
    p = ROOT / item['path']
    if not p.is_file() or hashlib.sha256(p.read_bytes()).hexdigest() != item['sha256']:
        discrepancies.append({'file': item['path'], 'expectedHash': item['sha256'], 'actualHash': hashlib.sha256(p.read_bytes()).hexdigest() if p.is_file() else None, 'expectedBytes': item['bytes'], 'actualBytes': p.stat().st_size if p.is_file() else None})
assert all(d['file'] == 'CODEX_PROMPT.md' for d in discrepancies), 'Affected handoff data must be reviewed'
incidents = read('incidents')
original = openpyxl.load_workbook(ROOT/'sources/baseline/incidents/Incident Database.xlsx', data_only=True)['Incident Database']
headers = [c.value for c in original[3]]
for incident in incidents:
    row = int(incident['id'].split('-')[-1])
    for col, header in enumerate(headers, 1):
        value = original.cell(row, col).value
        if hasattr(value, 'isoformat'): value = value.isoformat().split('T')[0]
        assert value == incident['values'][header], (row, header, value)
assert len(incidents) == 380
totals = {k: str(sum((Decimal(str(i['values'][k])) for i in incidents), Decimal(0))) for k in ['Downtime (hrs)', 'Act. Loss (k US$)', 'Pot. Loss (k US$)', 'Total Loss (k US$)']}
assert totals == {'Downtime (hrs)':'2261.1', 'Act. Loss (k US$)':'61886.46', 'Pot. Loss (k US$)':'5307.97', 'Total Loss (k US$)':'67194.43'}
decks = read('presentation_extraction')
assets = read('assets')
for asset in assets:
    wb = openpyxl.load_workbook(ROOT/asset['info_source']['file'], data_only=True)
    summary = wb['Performance Summary']
    asset['summary'] = [{'name':summary.cell(r,1).value, 'value':summary.cell(r,2).value, 'formula': next(k['value_or_formula'] for k in asset['performance_source_values'] if k['source']['cell'] == f'A{r}:C{r}'), 'basis':summary.cell(r,3).value, 'source':{'file':asset['info_source']['file'], 'sheet':'Performance Summary', 'cell':f'A{r}:C{r}'}} for r in range(4,17)]
    for c in asset['conditions']:
        assert [wb['Condition History'].cell(c['row'],n).value for n in range(3,7)] == c['measurements']
        assert wb['Condition History'].cell(c['row'],7).value == c['status']
    prod = openpyxl.load_workbook(ROOT/asset['production'][0]['source']['file'],data_only=True)['Sheet2']
    for record in asset['production']:
        row = int(record['source']['cell'].split(':')[0][1:])
        assert [prod.cell(row,c).value for c in range(2,9)] == list(record['values'].values())[1:]
    assert len(asset['production']) == 720 and len(asset['conditions']) == 26
    asset['report'] = next(d for d in decks if '/rca/' in d['file'] and asset['tag'] in d['file'])
    assert len(asset['report']['slides']) == 11
    link = next(i for i in incidents if i['id'] == asset['linked_incident_id'])
    assert link['values']['AR No.'] == asset['ar'] and link['values']['Tag Number'] == asset['tag']
    asset['plant'] = link['values']['Plant']
    asset['eventDate'] = link['values']['Date of Occur.']
version = 'caliber-v1-' + hashlib.sha256(''.join(s['sha256'] for s in inventory).encode()).hexdigest()[:12]
data = {'version':version, 'assets':assets, 'incidents':incidents, 'inventory':inventory, 'origins':read('source_origins'), 'metrics':read('verified_metrics'), 'official':read('official_text'), 'explanation':next(d for d in decks if '/explanation/' in d['file']), 'integrity':{'originalsChecked':22,'baselineFiles':16,'originalDiscrepancies':[], 'packageDiscrepancies':discrepancies, 'verifiedOn':'2026-10-02','originalRegisterTotals':totals}}
(OUT/'normalized.json').write_text(json.dumps(data,ensure_ascii=False,separators=(',',':')))
print(f"Normalized {len(assets)} assets and {len(incidents)} source rows; original values verified. {version}")
print('Package discrepancy retained: CODEX_PROMPT.md; all factual input files match manifest.')
